from pathlib import Path
import json,hashlib,shutil
import pandas as pd
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent;EXP=ROOT.parent/'experiment';A=ROOT/'assets'
E=json.loads((ROOT/'evidence.json').read_text(encoding='utf-8'))
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7);sec.left_margin=sec.right_margin=Cm(1.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 s=doc.styles[name];s.font.name='Calibri';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'等线');s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.18
doc.styles['Title'].font.size=Pt(25);doc.styles['Heading 1'].font.size=Pt(18);doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].paragraph_format.space_after=Pt(10)
for style in doc.styles:
 for el in list(style.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
foot=sec.footer.paragraphs[0];foot.alignment=2
run=foot.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
def p(t):return doc.add_paragraph(t)
def h(t):doc.add_heading(t,2)
def page(t):
 if len(doc.paragraphs)>1:doc.add_page_break()
 doc.add_heading(t,1)
def fig(n,caption,w=17.3):
 q=doc.add_paragraph();q.paragraph_format.space_after=Pt(3);q.add_run().add_picture(str(A/(n+'.png')),width=Cm(w));doc.add_paragraph(caption,'Caption')
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for i,row in enumerate(t.rows):
  for j,c in enumerate(row.cells):
   if widths:c.width=Cm(widths[j])
   tcpr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   tcpr.append(b);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6EF' if i==0 else ('F5F7F9' if i%2==0 else 'FFFFFF'));tcpr.append(sh)
   for q in c.paragraphs:
    q.paragraph_format.space_after=Pt(5);q.paragraph_format.space_before=Pt(5);q.paragraph_format.line_spacing=1.05
    for r in q.runs:r.font.size=Pt(9);r.bold=i==0
  if i==0:
   rpr=row._tr.get_or_add_trPr();rpr.append(OxmlElement('w:tblHeader'))
 p('')
def source(t):
 q=p('依据  '+t)
 for r in q.runs:r.font.size=Pt(8.5);r.font.color.rgb=RGBColor.from_string('555555')

doc.add_paragraph('以 DINOv2 为视觉骨干的\n秧苗自适应关键点算法技术拆解','Title')
p('从真实输入到自动教师 热图预测 结构图与表型计算')
p('以《2.docx》的 Stage 1—3 为起点  ｜  技术状态截至 2026年9月13日')
p('这套系统已经完成自动教师构建和学生检测器训练，能输出数量可变的点，并使点参与路径构建。当前还没有完成最终人工 GT 冻结，也没有证明学生比教师直接输出更适合表型测量。理解它，需要把视觉表征、自动监督、点检测和几何解码分开，再沿着同一张秧苗图把它们接起来。')
fig('06_overview','图1  当前候选链路。黄色为冻结骨干，粉色为可训练热图头及训练损失，绿色为规则与几何计算。虚线仅在训练中使用。这里的 G1′教师不由学生反向更新。')
p('阅读顺序：先校准原文中的术语，再看数据和 DINOv2 内部计算，然后追踪自动伪标签如何进入热图损失，最后解释可变点数、路径连接、局部短枝筛选和人工表型验证。所有秧苗图均来自现存 train/val 数据；彩色特征图来自真实张量，曲线来自冻结输出。')
source('S1—S5、S9—S12；完整路径与来源编号见末尾索引。')

page('1  从原说明出发校准六个概念')
p('《2.docx》已经抓住主线：早期候选器探索 → 自动教师监督学生 → 排除根须和修复短叶 → 比较 Teacher-direct 与 Student 的最终价值。下表保留这个叙述顺序，同时将容易混淆的地方对齐到代码与冻结证据。')
table(['原文理解','技术上的准确解释'],[
('“冻结 DINOV2 表征”','冻结整个 ViT-S/14-Reg 骨干的参数。reg 是 register token 版本标识，不是只冻结某个 reg 子模块。'),
('“对学生模型进行预训练”','DINOv2 已有外部自监督预训练；本项目进行的是冻结表征上的热图头训练。不是从零训练 DINOv2。'),
('“增加了教师的局部尺度 decoder”','教师增强改变训练目标；局部尺度路径 decoder 改变预测点之后的短枝接受规则。A/B/C/D 正是把二者拆开。'),
('“B=D>C>A”','B、D 在联合无错误数上同为22/40，但错误类型不同；不能写成普遍排名或统计等效。'),
('“人工复核之后才能进入训练”','多轮80轮训练已经完成。当前人工描迹用于评价和GT裁决，不进入热图监督；当前门槛是比较方法前冻结参考。'),
('“转为综述并比较历史固定关键点”','这是原说明中的方向设想，未成为当前授权路线。本技术文档继续解释现有自适应方案。')],[5.0,12.2])
h('三个名字分别指什么')
p('DINOv2 是视觉编码器。G1′是结合几何与视觉证据的自动候选器，经跨增强共识后提供离线伪标签。Student 是 DINOv2 骨干与卷积热图头组成的可学习检测器。G1′不通过梯度训练，Student 只更新热图头；目前没有学生反哺教师的迭代循环。')
p('“无人工关键点定义”允许规则先验存在。教师明确偏好端点、持续分叉和形状角点；但没有为每株指定固定编号、固定坐标和固定数量。这与“完全没有先验”不同，也不等于已经学会器官身份。')
source('S1《2.docx》；S2交接与状态；S3模型；S5教师；S12冻结配对报告。聊天只用于恢复疑问和语境，技术结论由代码与记录确认。')

page('2  Stage 1 到 Stage 3 的真实工作过程')
table(['阶段','已完成的工作','证据及其边界'],[
('G0与早期G1','数据质量审计；注意力峰、局部对比峰、单图聚类试探','注意力较稳定但覆盖不足；特征峰和聚类中心不等于表型支撑点。'),
('G1′与V3','几何—表征候选、九视图共识；87张教师；80轮训练','教师点数中位数9，学生初评点数5；尚不能说明路径语义正确。'),
('V4整株','构建300张整株图和三类自动掩膜；220张教师与训练','40张val初评重复性0.904；整株根须进入检测目标，暴露分析域问题。'),
('点条件图','学习点投影、骨架测地边、MST；零点和留一检查','修补旧骨架旁路；零学习点不再自动产生结构。'),
('地上部路线B','完整图留档；ROI教师216张；182/34内部划分训练','80轮；val重复性0.917；根须纯区域接受节点降为0。'),
('方案C与四组合','增强终端与过渡区目标；比较全局/局部路径解码','局部解码恢复真叶但引入假枝；增强教师未证明净收益。'),
('冻结与人工参考','冻结Student-B；D只保留对照；三轮人工描迹及可靠性','第二测量者两阶段均完成，当前等待双人共识、必要重描和最终GT。')],[2.8,7.0,7.4])
p('V1—V4 是数据版本；G0、G1、G1′是工作阶段或候选器代号；A/B/C/D 是教师来源与路径解码器的二维组合。它们不能串成“模型升级到V4，再升级到D”的单一序列。')
p('V3 的 test 已被历史初评接触，其98张在V4归入训练开发来源。V4新val/test由新的扫描采集组构成；同帧兄弟裁剪和同采集组不跨划分。当前采用开发val的人工证据不能当作最终泛化性能。')
source('S2；S8数据报告；S9训练摘要；S10—S12模块和人工复核报告。')

page('3  真实输入及地上部有效域')
fig('01_input','图2  v4_val_0004 的真实输入演变。第一行是扫描裁剪复核图，第二行是白底标准化整株图，第三行是从整株及自动器官掩膜派生的地上部分析输入。处理图不冒充原始采集图。')
h('为什么完整保存后还要做 ROI')
p('整株图保存颖果、叶鞘和可见根，便于追溯。地上部表型却不应从米粒或根须起算。旧路线的40张审核中，28张基部被判wrong，38条备注涉及根系干扰。若这些组织继续产生教师正点，热图头就会认真学习它们，后处理很难彻底纠正。')
p('现有程序利用shoot、seed_base_root和full_plant自动掩膜，定位颖果邻近的shoot边界，并从远端绿色组织估计茎叶延伸方向。按方向保留地上部及窄基部过渡带，生成 phenotype_roi 与 basal_transition。方向来自当前图像，不能简化成固定“保留图像上半部”。')
p('训练及推理将有效域外像素置白，再送入视觉骨干；图重建使用相同分析域的骨架，基部候选须接近shoot侧过渡区。自动过渡区只是区域约束，最终起算节点仍必须来自检测点。当前热图损失代码覆盖整个输出画布，没有额外逐像素ROI损失掩码；域外主要通过白底输入和无正点目标受到抑制。')
source('S7 phenotype_roi_basal_anchor.py；S8；S10有效域报告；S4 PseudoPointDataset与heatmap_loss。')

page('4  图像怎样变成 DINOv2 的 token')
fig('07_backbone','图3  按本地官方实现重绘的 ViT-S/14-Reg 机制。主干共12层、每层6个注意力头；浅黄色模块参数全部冻结。下方展开一个包含两次残差相加的Transformer块。')
p('输入先等比例缩放并补白为518×518，不拉伸苗体。以14×14不重叠卷积切块后，网格为37×37，共1369个patch。每个patch投影到384维。加入CLS token、位置编码及4个register token后，Transformer序列长度为1374。')
p('每个注意力头的维度为384÷6=64。先对token做LayerNorm，再线性生成Q、K、V，以softmax(QKᵀ/√64)V混合各位置的信息。6个头拼接并投影后经LayerScale与残差相加；另一支经LayerNorm、384→1536→384的MLP及残差相加。层叠后，一个叶尖patch的表示可以含有主轴和其他叶片提供的上下文。')
p('输出经过最终LayerNorm，拆成CLS、register和patch tokens。检测头只取 x_norm_patchtokens，去掉5个非空间token，将1369×384重新排列成384×37×37。不能将register画成叶片上的4个关键点，也不能将CLS分类向量直接当作定位热图。')
h('DINOv2 的外部预训练与本项目训练')
p('DINOv2通过大规模无人工标签的视觉预训练获得通用表征；本项目加载现成权重，未复现它的大规模预训练。DINOv2自身的自监督教师—学生体系与本项目“G1′伪标签→热图Student”属于不同训练层级，不能把其预训练损失画进本项目热图优化链路。')
source('S3；S6官方vision_transformer.py、attention.py、block.py；外部原始文献R1。')

page('5  register 的作用及冻结边界')
p('reg 来自register。相关研究观察到，ViT可能把低信息背景patch用作内部计算空间，产生高范数伪影。额外register tokens为这种内部计算提供独立位置，有助于获得更平滑的稠密特征和注意力图。它们不携带固定图像坐标，也不表示预定义器官。该动机解释了版本选择，不能替代本秧苗数据上的有无register消融。')
table(['模块或量','当前是否更新','具体含义'],[
('14×14 patch embedding','否','卷积投影权重、偏置固定'),('CLS、位置编码、4个register参数','否','可学习参数数值固定；每次前向产生的token激活仍随图像改变'),('12个Transformer块','否','QKV、输出投影、MLP、LayerNorm、LayerScale参数均不优化'),('最终LayerNorm及其他骨干参数','否','configure_backbone遍历所有backbone.parameters()'),('卷积热图头','是','3个卷积与2个GroupNorm，共830,401参数'),('概率阈值、NMS、几何路径解码','不是梯度参数','按已锁定配置运行，不由反向传播学习')],[6.4,2.2,8.6])
h('为什么先冻结整个骨干')
p('初步实验首先问“自动教师目标能否被一个轻量头学习”。冻结预训练表征可减少优化参数，让这个问题更容易解释，也适合当前较小的训练集。骨干不再随噪声伪标签改变，便于将检测头拟合问题与表征适配问题区分。这是工程选择及其合理动机，不是已经证实冻结优于微调。')
p('实际配置 freeze_backbone=true、unfreeze_last_blocks=0。优化器只收集requires_grad为真的参数；配置中虽然有backbone_learning_rate=10⁻⁵，当前并不会创建骨干更新组。解冻最后两层仅有计划配置，尚未完成正式消融，不能写成已使用模块。')
p('“冻结”不等于关闭前向计算，也不自动等于eval模式。训练入口调用model.train()；是否存在随机行为还取决于骨干具体配置。可靠判据是参数requires_grad和优化器分组。本次展示使用eval与inference_mode，没有任何参数更新。')
source('S3 configure_backbone；S4优化器分组；S9配置；register机制见原始文献R2。')

page('6  真实特征图并不是预测热图')
fig('02_features','图4  v4_val_0004 的真实 DINOv2 输出。PCA仅把384维投影为显示用RGB颜色；CLS注意力和局部对比均由真实张量计算。三者均不是人工关键点标签，也不是学生概率热图。')
h('教师与学生读取不同的表征')
p('教师调用 get_intermediate_layers(n=4)，逐层对patch向量L2归一化，平均最后4层后再次归一化，得到last4avg。局部对比先取3×3邻域平均特征，再计算1−cos(f，邻域均值)。相邻位置外观明显变化时，对比度可能升高，但背景边界或伪影也可能很高。')
p('教师另通过最后一个注意力块的QKV计算CLS对全部token的softmax权重，跨6个头平均，再去掉CLS与register位置，得到37×37的patch注意力。它度量相对注意分配，不是“这里是关键点”的校准概率。若代码无法抓取QKV，会明确标为CLS—patch余弦相似度回退；本次实际提取得到exact_last_block_cls_attention。')
p('学生调用forward_features，只读取最后一层的x_norm_patchtokens。之后由有监督的卷积头把这些特征转换为点响应。该头没有直接拼接教师的注意力图、骨架图或last4avg；几何规则通过伪标签影响学生，而不是作为额外输入通道进入热图头。')
h('早期四候选器到底检验什么')
p('注意力峰检验显著性，局部对比峰检验视觉差异，HDBSCAN检验特征簇的代表中心，G1′检验结构候选与表征的融合。前三类是直接找点的无训练基线，不是当前学生的三层网络。早期比较支持继续探索G1′，不能推出其在最终表型上已经优于所有替代方法。')
source('S5；S6 g1_dinov2_feasibility.py中的extract_representations和feature_local_contrast；S13四类摘要。')

page('7  G1′如何提出数量可变的候选')
fig('03_teacher',f"图5  真实训练样本 {E['teacher_sample']['dataset_id']}。骨架由现有输入和规则计算；点坐标、共识置信度来自已归档教师JSONL；最右图按训练代码构造实际高斯目标，未用验证图生成教师。")
p('教师先从标准化RGB估计高置信前景，经形态学处理、细化与小连通体过滤得到骨架。8邻域度数为1的像素形成端点，度数≥3的区域经过聚合与支臂持久性检查后形成分叉候选。形状角点来自前景轮廓响应；DINOv2局部显著点须靠近骨架且处于前景内对比度的高分位区域。')
p('每个候选的融合分数为 s = prior + 0.24×local_contrast + 0.08×attention。端点prior=0.68，持续分叉=0.76，形状角点=0.32，特征峰=0.30。端点/分叉直接进入结构候选，角点和特征峰另有分数门槛。按分数排序，用max(8，round(0.05D))的空间间距去重；D为模型画布上前景包围盒对角线。')
p('因此，G1′中几何先验有实质作用。DINOv2补充评分和局部证据，不宜把系统描述为“骨干独立自主发现全部关键点”。当前教师生成默认安全上限为30；早期候选器对比曾使用20。上限是工程截断保护，不能混同学生64点上限，更不能据此声称固定点数。')
p('原路线B教师使用普通候选去重。后来结构覆盖增强教师会优先保留不同终端，并额外提出shoot侧基部过渡候选，以避免小叶端点被主轴邻点压制。这是历史方案C/D中的教师改动，当前冻结主候选B没有把它作为必要组成。')
source('S5 structural_candidates、add_candidate、consensus；S12冻结组合。')

page('8  九视图共识怎样成为训练目标')
p('一张train图生成identity、水平翻转、旋转+10°、亮度×0.75、旋转−10°、缩放0.9、缩放1.1、平移以及对比度×1.25，共9个视图。教师分别提候选；对几何增强点做逆变换，使所有坐标回到identity画布。光度增强坐标不动。')
p('匹配使用Hungarian一对一分配，之后按距离≤0.05D过滤。以identity候选为参照，统计在多少视图找到对应点。它不是对所有视图候选做无条件并集；identity中没有的点不会仅凭其他视图出现就自动成为基础目标。')
table(['共识量','实际计算或门槛'],[
('presence r','被观测视图数÷9；r≥0.75，即至少7个视图'),
('定位误差 e','其余匹配视图逆变换距离/D的中位数；e≤0.025'),
('平均教师分数 s̄','匹配观测候选分数的均值，置信度计算时截断至[0,1]'),
('共识置信度 c','r×exp(−e/0.025)×clip(s̄,0,1)；c≥0.35'),
('可训练图','接受点不少于2，并且identity候选没有触发教师安全上限')],[4.8,12.4])
p('一个点反复出现，只能说明它对这些增强稳定。根须或错枝也可能稳定出现，所以共识不是语义真值；先约束分析域、再检查真实器官路径，仍然必要。')
h('从离散点到单通道软目标')
p('将每个伪标签坐标投影到129×129输出网格，以其共识置信度为峰高，生成σ=1.6输出像素的高斯。多个高斯逐像素取最大值，不做相加：T(u)=maxᵢ[cᵢ exp(−‖u−uᵢ‖²/(2σ²))]。局部核仅写入约3σ范围，背景为0。')
p('这种监督允许同一通道出现多个峰，且峰数随当前图的稳定结构变化。它不需要按“第一叶尖、第二叶尖、第三叶尖”开设固定通道。教师坐标包含kind供审计，但学生并不学习kind分类。')
source('S5 make_transforms与consensus；S4 gaussian_heatmap。')

page('9  可训练热图 decoder 的逐层计算')
fig('08_head','图6  当前学生卷积头的精确结构。卷积均含偏置；3×3卷积padding=1。GroupNorm为8组。先在37×37上预测logits，再双线性插值到129×129，最后在解码时sigmoid。')
table(['步骤','输出张量 单图','可训练参数'],[
('最后一层patch重排','384×37×37','骨干冻结'),('Conv 3×3 384→192','192×37×37','663,744'),('GroupNorm 8组 + GELU','192×37×37','384'),('Conv 3×3 192→96','96×37×37','165,984'),('GroupNorm 8组 + GELU','96×37×37','192'),('Conv 1×1 96→1','1×37×37','97'),('双线性插值 align_corners=false','1×129×129','0'),('总计','共享单通道点响应','830,401')],[7,5,5.2])
p('3×3卷积结合邻近patch的语义，GroupNorm按通道组归一化，适合小批量训练，GELU提供非线性。第二层压缩通道，最后1×1将96维局部描述压为一个logit。因而这是一个轻量密集预测头，不是为每个点设置query的Transformer decoder。')
p('518÷4向下取整得到129，而不是129.5。上采样增加输出采样位置，不会凭空恢复14像素patch内的全部细节。细短叶在输入縮小、patch表征、热图峰融合和后续几何筛选中都可能丢失，不能把所有漏叶归为同一种原因。')
p('参考图片采用多阶段卷积残差、两张热图、深度掩膜和TAPIR追踪。现有秧苗实现没有这些模块。图形的配色、块级展开和输入输出对应方式可参考它，机制必须遵循本项目实际代码。')
source('S3 AdaptivePointDetector.__init__与forward；参数量由本次实例化核对。')

page('10  学生到底优化了什么')
p('热图logit为z，概率P=sigmoid(z)。主损失是在整个129×129画布上求均值的加权二元交叉熵：Lheat=mean[(1+80T)·BCEWithLogits(z,T)]。正目标附近权重更大，防止大量背景像素压倒少量关键点峰。教师置信度既决定目标峰高，又影响局部权重。')
p('数量辅助项先估计 N̂=ΣP/(2πσ²)，再计算SmoothL1(N̂,Nteacher)，总损失 L=Lheat+0.02Lcount。它用高斯总质量近似点数，不是专门的计数输出头。由于目标峰高并非全为1，重叠峰取最大值且边界有截断，质量近似不严格等于真实峰数；数量项只能作为辅助约束。')
fig('09_training','图7  路线B的真实80轮训练日志。纵轴为自动教师拟合损失，不是表型误差。最佳内部验证轮次和损失由history.csv计算。')
p('训练采用AdamW，头学习率0.001，weight_decay=0.01；batch_size=4、梯度累积2，完整累积时有效批量约8。训练增强为0.5概率水平翻转和0.85—1.15亮度缩放，坐标随翻转同步变换。RTX3090训练启用AMP；配置没有额外学习率调度器。')
p('路线B从220张train中按质量门槛排除4张，216张教师样本按种子和ID哈希分成182张梯度训练、34张内部验证。80轮保存的global_steps=3680是批次计数；每轮46批、约23次优化器更新，不能把它直接称为3680次参数更新。最佳验证损失0.2845，总训练计时约104秒，不含教师生成和人工审计。')
source('S4训练入口、deterministic_split、heatmap_loss；S9 history.csv、training_summary.json。')

page('11  概率热图如何变成可变数量坐标')
fig('04_heatmap','图8  v4_val_0004 的真实Student前向。中间为概率，不是PCA或教师高斯；右侧为实际阈值和NMS得到的3个点。置信度展示使用本次CPU数值。')
p('解码先对概率做5×5最大池化。一个位置同时满足“等于邻域最大值”和“概率≥0.35”才成为峰。按置信度排序，超过64个时才截断到安全上限；fixed_k=0，不要求输出固定K。模型可能输出0、3或更多个点，不通过补点保证条数。')
table(['点','模型画布 x y','CPU置信度'],[(f'p{i+1}',f"{q['x']:.4f}, {q['y']:.4f}",f"{q['confidence']:.6f}") for i,q in enumerate(E['samples']['v4_val_0004']['points'])],[2,10.2,5])
p('网格坐标(u,v)映射到518画布时，代码使用x=u×517/128，y=v×517/128。随后减去padding并除以等比例scale，恢复标准化源图坐标。示例源图1616×339，scale=0.3205445545，pad_y=204；第一点在源图约为(189.009，119.621)。这些坐标是标准化图像坐标，不是原始TIFF整帧坐标。')
p('本次3张CPU前向分别输出3、3、5点；与历史GPU保存点在模型画布坐标上的最大差约5.7×10⁻¹⁴ px，可视为数值精度内一致。置信度略有CPU/GPU差异，因此后文路径和人工结果继续引用原冻结输出，不把本次数值替换成新实验结果。')
source('S3 decode；S14本次evidence.json；S9保存points.csv。')

page('12  点条件结构图怎样控制路径')
p('早期程序从完整骨架的端点和分叉独立生成路径，学习点只补充样条支撑。这意味着删掉全部学习点也可能有路径；当时graph_success=1不能证明自适应点有效。当前点条件图专门切断这个旁路。')
h('投影  接受  连边')
p('对预测点pᵢ找最近骨架像素qᵢ，仅在‖pᵢ−qᵢ‖≤0.025D时接受。多个点投到同一像素则保留最高置信度。接受点成为唯一图节点；骨架端点、分叉不自动补成新节点。这里0.025来自此前val扫描，并已锁定。')
p('骨架提供8邻域测地路由，同一连通分量中计算节点两两最短路径，再构造最小生成树；一条抽象图边具有具体的骨架像素路径。节点少于2时没有测地边。覆盖来自这些边的空间并集，不来自把整张骨架直接当作输出。')
h('为什么 MST 的度数不能直接当叶尖')
p('两条测地边可能在空间上复用或经过同一个真实终端。某节点在抽象MST中度数为2，不保证它在空间几何上处于中间。旧实现曾造成39/40张只有1条路径；v4_val_0035暴露此问题后，改为检查点条件边并集的空间终端，并要求终端有学习节点支撑。')
table(['当前路线B的40张val','冻结数值'],[
('输入预测点 / 接受图节点','161 / 143'),('接受节点数中位数','3'),('骨架覆盖中位数 / 均值','0.9626 / 0.9068'),('覆盖低于0.90','5 / 40'),('零点对照覆盖','0'),('留一删除后覆盖下降的节点','124 / 143')],[11,6.2])
p('留一贡献ΔCᵢ=C(P)−C(P去掉pᵢ)检验点是否影响覆盖。正贡献说明点控制结构，零贡献可能是同路径冗余采样，不能直接证明该点对曲率或长度无用。骨架覆盖是自动结构代理，错连也可能有高覆盖，仍须人工语义和表型参考。')
source('S10 point_conditioned_graph.py及graph_v2_phenotype_roi_val/summary.json。')

page('13  几何 decoder 为什么能恢复短叶')
p('首先在已有图节点中选择基部：节点须位于带容差的ROI内，距离shoot侧过渡区≤max(12，0.08D)，按距离优先、置信度次优排序。无合格节点就明确失败，不用自动区域中心或根须补一个“基点”。')
p('从基部向可达终端计算路径，以测地距离最远者作为主路径，其他终端在已恢复路径树的最后共享位置分出侧枝。路径数来自终端与筛选，而非提前规定叶片数。这里的主路径是算法几何选择，不自动等价于植物学主茎或人工主参考叶。')
h('全局规则与局部叶宽规则')
p('全局规则要求侧枝长度≥max(8，0.04D)，整株很长时容易删除真实短叶。局部规则在分支远端一半取ROI距离变换的正半径中位数，叶宽w=2×median(radius)，要求分支长度≥max(4，1.5w)，并且终端置信度≥0.5×当前图节点置信度中位数。使用远端避免粗的连接部位支配细叶阈值。')
table(['真实样本或量','数值与作用'],[
('v4_val_0004 全局阈值','13.8971模型像素'),('同一短枝长度','10.4853模型像素，低于全局阈值'),('局部叶宽与阈值','w=5.6000；max(4,1.5w)=8.4000'),('终端置信度与门槛','0.553594 ≥ 0.313815'),('最终决定','局部规则接受；人工A→B配对记录为恢复1片真叶'),('v4_val_0006 的另一分支','4.4142 < 局部阈值6.5907，显式拒绝')],[7.5,9.7])
p('这个decoder是距离变换、路径长度和置信度的确定性规则，没有卷积层、训练参数或反向传播。它只能决定现有终端对应的短枝是否保留，不能从没有学习点的地方新发现一片叶。局部放宽也会接受伪枝；这正是人工对照中新增假枝和错连的来源之一。')
source('S11 _choose_learned_phenotype_base、_local_branch_scale；冻结terminal_branch_decisions.jsonl；S12配对审核。')

page('14  同一条链中的真实成功与失败')
fig('05_cases','图9  三个真实val案例，左为模型输入，中为冻结接受节点，右为存档的基部到终端路径。0004体现短叶恢复；0002因缺少合格学习基部而零路径；0034属于局部解码引入假枝/错连的人工诊断样本。此图不展示或修改人工GT。',17.2)
p('0004：模型输出3点且全部被接受，骨架覆盖约0.9694；局部decoder输出2条候选路径。前页数值说明较短分支如何通过局部阈值，人工配对结果确认恢复真叶。0002：虽然有3个检测点，但只接受2个，且没有满足过渡区约束的基部，零路径是有意义的失败输出。0034：5个预测点并不等于5片真实叶，必须检查共享路径与新增枝条的语义。')
source('S9 points.csv；S10 nodes.csv；S11 paths.jsonl、per_image.csv；S12人工配对结果。')

page('15  样条与表型计算的单位和定义')
p('已有路径先用RDP选取几何转折支撑，再并入路径附近的学习点支撑，用PCHIP做坐标插值。几何支撑点用于数值拟合，不能生成缺失的器官；它们也不是新增的学习检测点。PCHIP减少普通三次样条在急弯处越出骨架的过冲，但不保证所有曲线都位于真实叶片中心。')
h('历史候选 CSV 与正式表型协议应分开读')
p('旧candidate_phenotypes.csv保存的侧枝长度通常是分出位置到终端的侧枝段；列名spline_length_px的像素处于518模型画布尺度。它不能直接改名为正式“基部—叶尖结构路径长度”。例如v4_val_0001的主路径样条长351.0916，侧枝段17.2627，后者不是从共享地上部基点出发的完整长度。')
p('正式phenotype-geometry-v1从完整基部到叶尖路径出发：去除连续重复点，按累计弧长参数化，逐坐标PCHIP并固定重采样240点。长度由相邻采样点欧氏距离累加；弦长为首尾距离；弯曲度可用路径长/弦长。平均绝对曲率是探索指标，受描迹与平滑方式影响较大。')
p('分化角使用共享段容差max(2 px，0.01D)，在分开后的0.05D局部区间估计方向，与reference_main_path比较。它不是整条叶片拟合直线与图像竖直方向的夹角。参考主路径按结构路径长度、弦长、极角、trace_uuid依次确定，规则在测量前锁定。')
table(['坐标或单位','怎样使用'],[
('518模型画布像素','图构建、局部短枝阈值和历史候选CSV；不能直接套600 dpi'),('标准化源图像素','去padding并除scale；须核对标准化过程相对原扫描的缩放'),('bbox归一化','路径长/D；分子分母必须在同一坐标尺度'),('扫描元数据换算毫米','按扫描DPI及标准化缩放链换算；标注derived from scanner metadata'),('正式物理精度','仍须已知尺寸实物标尺校验，元数据换算不是精度认证')],[5.3,11.9])
source('S11路径保存与候选CSV；S15 phenotype_gt_geometry.py和人工轨迹计算锁定规范。本文未运行任何新的方法表型比较。')

page('16  A B C D 对照真正支持什么')
table(['组合','教师目标','路径decoder','自动路径总数'],[('A','原路线B','全局整株尺度','64'),('B 当前冻结候选','原路线B','局部叶宽尺度','76'),('C','结构覆盖增强','全局整株尺度','63'),('D 冻结对照','结构覆盖增强','局部叶宽尺度','74')],[4.2,4.5,5.2,3.3])
p('这里A/B用同一个学生checkpoint，C/D用另一个学生checkpoint；换decoder不重训热图头。A→B与C→D检验解码规则，A→C与B→D检验教师目标改变后训练出的学生。四组合不是Teacher-direct与Student的直接比较。')
table(['40张配对人工评价','A','B','C','D'],[
('真实路径总数',63,70,63,69),('漏叶总数',23,16,22,16),('假枝总数',0,4,0,3),('错连图像数',6,10,3,5),('基部错误图像数',5,5,7,7),('联合无错误图像数',18,22,17,22)],[7.6,2.4,2.4,2.4,2.4])
p('局部decoder在两套教师下分别净恢复7、6片真叶，且没有真叶丢失；但新增4、3条假枝以及4、2张错连。它有稳定覆盖收益，也有明确精度代价。')
p('增强教师在全局decoder下人工偏好5:5，联合无错误从18降到17；在局部decoder下，B/D同为22张联合无错误、16片漏叶。D减少错连，却增加基部错误并少1条真实路径。由此冻结B为当前最小充分Student候选，停止修补C/D；不能概括为增强教师更优或B与D统计等效。')
p('下一比较必须让Teacher-direct与Student-B共用完全相同的冻结graph、局部decoder和表型算子，只替换输入点来源。只有这样才可能判断热图Student在误差、稳定性或速度上是否提供额外价值。当前尚无这项表型比较结果。')
source('S12 A_B_C_D人工配对复核与因果分析_20260803.md；S15测量前冻结协议。')

page('17  人工参考的作用和当前证据上限')
p('无人工关键点监督并不排斥独立人工评价。过去的val审核用于发现通用失败模式并选择方法，属于开发过程；现在的模型盲描迹只建立评价参考，不能写回教师目标、热图损失或逐图规则。部署推理链不包含人工逐株点选。当前训练已完成，等待的是公平评价所需的参考可靠性。')
table(['人工参考环节','截至2026年9月13日的状态'],[
('样本','12张Core＋4张Diagnostic；两层统计分别报告'),('三轮描迹','Rater1两轮35/36条；Rater2一轮30条，均已归档'),('重复测量误差','长度相对MDC95取较大者3.72%；分化角MDC95为17.08°'),('初步分歧','13/16株触发预锁定裁决条件；类别有重叠'),('两阶段裁决','Rater2两阶段13图、30结构均提交；28结构选候选，1个需重描、1个保留争议'),('剩余门槛','7张优先讨论；全部13张逐结构共识、必要重描或第三方裁决及最终GT冻结'),('随后允许比较','按冻结协议运行Teacher-direct/B/D；正式消融、多种子和V4 test后置')],[5.2,12])
p('MDC95描述人工重复差异的尺度，不是算法准确率，也不保证每个裁剪端点都是真实叶尖。“全部标可测”是测量者意见，仍需针对图像完整性和叶片身份取得共识。旧13张A/B/C选择只是一位测量者的独立裁决，不能当成两位专家一致。')
h('目前可以写到什么程度')
p('可以说：自动结构—表征共识目标可以监督一个冻结DINOv2上的轻量热图头；点数随峰值动态变化；接受点能够控制结构覆盖；地上部有效域减轻根须干扰；局部尺度解码恢复部分短叶。')
p('尚不能说：DINOv2单独发现了植物学关键点；Student已经超过Teacher-direct；增强教师有净收益；表型长度与角度已经达到指定精度；当前val成绩代表最终独立测试性能。当前真正需要补齐的是点、路径、表型之间的证据链，而不是仅增加训练轮数。')
source('S2最新状态；S15三轮可靠性报告及第二阶段几何核验与双人共识导航。')

page('18  代码与数据来源索引')
p('项目根目录为 D:/kp/adaptive_keypoint_discovery_reboot/。以下相对路径均从该根目录解析；本地审计起点为b0dc39f，交付前同步befeda6的裁决更新。正文提及的历史结论按其记录日期解释，最新门槛以2026年9月13日状态为准。')
refs=[
('S1','C:/Users/F/Desktop/2.docx','原始Stage 1—3说明；未覆盖原文件。'),
('S2','AGENTS.md；ACCOUNT_HANDOFF.md；PROJECT_STATE.md','研究边界、历史演化与最新裁决门槛。'),
('S3','experiment/adaptive_point_model.py','AdaptivePointDetector：冻结、卷积头、forward及decode。'),
('S4','experiment/train_adaptive_point_detector.py','数据缩放、gaussian_heatmap、heatmap_loss、AdamW与训练循环。'),
('S5','experiment/generate_g1prime_pseudolabels.py；g1_prime_structural_support.py','候选评分、九视图、逆变换、共识筛选。第二文件同在experiment/。'),
('S6','experiment/g1_dinov2_feasibility.py；experiment/third_party/dinov2_git/dinov2/','表征提取与官方models/vision_transformer.py、layers/attention.py、block.py。'),
('S7','experiment/phenotype_roi_basal_anchor.py','derive_phenotype_roi、load_phenotype_input。'),
('S8','experiment/data_stage_clean_v4_fullplant_candidate/','manifests/train.csv、val.csv；images/val；raw_review/val；三类masks。'),
('S9','experiment/training_outputs/core_dinov2_v4_phenotype_roi/','best.pt、history.csv、training_summary.json；检测输出在evaluation_outputs/core_dinov2_v4_phenotype_roi_val/。'),
('S10','experiment/point_conditioned_graph.py','图算法；evaluation_outputs/point_conditioned_graph_v2_phenotype_roi_val/中的summary.json、nodes.csv。'),
('S11','experiment/point_conditioned_organ_paths.py','路径算法；evaluation_outputs/point_conditioned_organ_paths_v2_phenotype_roi_local_decoder_val/中的paths、per_image、candidate_phenotypes及terminal_branch_decisions。'),
]
for code,path,desc in refs:
 q=p(code+'  '+desc);q.paragraph_format.space_after=Pt(3)
 for r in q.runs:r.bold=True;r.font.size=Pt(9.5)
 q=p(path);q.paragraph_format.space_after=Pt(7)
 for r in q.runs:r.font.size=Pt(9)

page('19  文献与可追溯附件')
for code,path,desc in [
('S12','experiment/A_B_C_D人工配对复核与因果分析_20260803.md','冻结组合、人工计数与失败代价；对应factorized_method_review_completed.csv。'),
('S13','experiment/四类方法关键技术摘要.md','早期无训练候选器的区别。'),
('S14','technical_breakdown_20260913/evidence.json；assets/','本次前向审计、真实特征/热图NPZ、PNG和SVG机制图；build_evidence.py、build_diagrams.py可重建图件。'),
('S15','experiment/phenotype_pilot_protocol/','phenotype_gt_geometry.py、人工跨轮次匹配与轨迹计算锁定规范_20260803.md、三轮人工GT匿名匹配与可靠性初步结果_20260905.md、第二阶段几何核验与双人共识导航_20260913.md。')]:
 h(code+'  '+desc);p(path)
h('原始文献')
p('R1  Oquab et al. DINOv2: Learning Robust Visual Features without Supervision. arXiv:2304.07193. https://arxiv.org/abs/2304.07193')
p('R2  Darcet et al. Vision Transformers Need Registers. arXiv:2309.16588. https://arxiv.org/abs/2309.16588')
p('官方实现  https://github.com/facebookresearch/dinov2')
p('项目镜像  https://github.com/KaifengWei/adaptive-keypoint-discovery')
h('本次复核的范围')
p('原始说明的疑问与已有聊天记录相互核对，正文机制以本地源码、配置、原始数值文件为准。历史会话中“教师是否反向学习”“人工审核是否进入训练”的讨论，分别落实为第1、8、17节。旧会话的阶段性建议不覆盖现行冻结边界。')
p('图件来源：v4_val_0004、0002、0034用于既有Student的少样本CPU前向及存档输出说明；另一个train样本用于展示已存在伪标签与高斯目标。未进行参数更新、阈值扫描、GT解盲或新方法表型比较；V4 test读取数为0。PCA显示色彩没有器官类别含义。')
p('冻结checkpoint的SHA-256：'+E['checkpoint_sha256'])
p('文档保留了当前已证实的能力及失败边界。后续论文方法描述应与最终GT冻结和方法比较结果同步，不把这份技术解释当作尚未完成实验的替代证据。')
source('真实张量和历史CSV可通过S14审计索引逐一追溯。')

out=ROOT/'DINOv2自适应关键点算法技术拆解.docx';doc.save(out)
shutil.copy2(out,Path('C:/Users/F/Desktop')/out.name)
print(out)
